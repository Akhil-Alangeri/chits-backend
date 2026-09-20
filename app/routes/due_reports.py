from datetime import date
from typing import Optional

from fastapi import APIRouter, Depends, HTTPException, Query, status
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from io import BytesIO

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.enums import TA_LEFT
from reportlab.platypus import (
    SimpleDocTemplate,
    Table,
    TableStyle,
    Paragraph,
    Spacer,
)

from app.core.database import get_db
from app.models.admin.user import User
from app.schemas.due_report import DueReportResponse
from app.services.due_reports_service import (
    get_due_report,
    get_due_report_export_data,
)

from app.core.dependencies import get_current_user


router = APIRouter(
    prefix="/reports",
    tags=["Reports"],
)


# ============================================================
# Helpers
# ============================================================

def _validate_date_range(
    date_from: Optional[date],
    date_to: Optional[date],
):
    if date_from and date_to and date_from > date_to:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="date_from cannot be later than date_to.",
        )


def _handle_company_error(exc: ValueError):
    if str(exc) == "Company not found.":
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Company not found.",
        )

    raise exc


# ============================================================
# Due Report
# ============================================================

@router.get(
    "/due",
    response_model=DueReportResponse,
)
def due_report(
    date_from: Optional[date] = Query(None),
    date_to: Optional[date] = Query(None),

    chit_id: Optional[int] = Query(
        None,
        gt=0,
    ),

    member_id: Optional[int] = Query(
        None,
        gt=0,
    ),

    report_status: str = Query(
        "all",
        alias="status",
        pattern="^(all|due|partial|overdue)$",
    ),

    search: Optional[str] = Query(None),

    page: int = Query(
        1,
        ge=1,
    ),

    page_size: int = Query(
        10,
        ge=1,
        le=100,
    ),

    db: Session = Depends(get_db),

    current_user: User = Depends(
        get_current_user
    ),
):
    # ---------------------------------------------------------
    # Company scope
    # ---------------------------------------------------------

    if not current_user.company_id:
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail="User is not associated with a company.",
        )

    # ---------------------------------------------------------
    # Date validation
    # ---------------------------------------------------------

    _validate_date_range(
        date_from,
        date_to,
    )

    # ---------------------------------------------------------
    # Service
    # ---------------------------------------------------------

    try:
        return get_due_report(
            db=db,
            company_id=current_user.company_id,

            date_from=date_from,
            date_to=date_to,

            chit_id=chit_id,
            member_id=member_id,

            report_status=report_status,

            search=search,

            page=page,
            page_size=page_size,
        )

    except ValueError as exc:
        _handle_company_error(exc)


# ============================================================
# Excel Export
# ============================================================

@router.get(
    "/due/export/excel"
)
def export_due_report_excel(
    date_from: Optional[date] = Query(None),
    date_to: Optional[date] = Query(None),

    chit_id: Optional[int] = Query(
        None,
        gt=0,
    ),

    member_id: Optional[int] = Query(
        None,
        gt=0,
    ),

    report_status: str = Query(
        "all",
        alias="status",
        pattern="^(all|due|partial|overdue)$",
    ),

    search: Optional[str] = Query(None),

    db: Session = Depends(get_db),

    current_user: User = Depends(
        get_current_user
    ),
):
    if not current_user.company_id:
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail="User is not associated with a company.",
        )

    _validate_date_range(
        date_from,
        date_to,
    )

    try:
        report = get_due_report_export_data(
            db=db,
            company_id=current_user.company_id,

            date_from=date_from,
            date_to=date_to,

            chit_id=chit_id,
            member_id=member_id,

            report_status=report_status,

            search=search,
        )

    except ValueError as exc:
        _handle_company_error(exc)

    # ---------------------------------------------------------
    # Workbook
    # ---------------------------------------------------------

    workbook = Workbook()

    worksheet = workbook.active
    worksheet.title = "Due Report"

    company_name = report["company_name"]
    company_id = report["company_id"]

    # ---------------------------------------------------------
    # Header
    # ---------------------------------------------------------

    worksheet["A1"] = company_name
    worksheet["A1"].font = Font(
        bold=True,
        size=16,
    )

    worksheet["A2"] = (
        f"Company ID: {company_id}"
    )

    worksheet["A4"] = "DUE REPORT"
    worksheet["A4"].font = Font(
        bold=True,
        size=14,
    )

    # ---------------------------------------------------------
    # Summary
    # ---------------------------------------------------------

    summary = report["summary"]

    worksheet["A6"] = "Total Due"
    worksheet["B6"] = summary["total_due"]

    worksheet["A7"] = "Total Paid"
    worksheet["B7"] = summary["total_paid"]

    worksheet["A8"] = "Total Remaining"
    worksheet["B8"] = summary["total_remaining"]

    worksheet["A9"] = "Total Records"
    worksheet["B9"] = summary["total_records"]

    worksheet["D6"] = "Due"
    worksheet["E6"] = summary["due_records"]

    worksheet["D7"] = "Partial"
    worksheet["E7"] = summary["partial_records"]

    worksheet["D8"] = "Overdue"
    worksheet["E8"] = summary["overdue_records"]

    worksheet["D9"] = "Unique Members"
    worksheet["E9"] = summary["unique_members"]

    # ---------------------------------------------------------
    # Table headers
    # ---------------------------------------------------------

    header_row = 11

    headers = [
        "Installment",
        "Due Date",
        "Chit",
        "Member No",
        "Member Name",
        "Customer Code",
        "Due Amount",
        "Paid Amount",
        "Remaining",
        "Status",
        "Days Overdue",
    ]

    for column_index, header in enumerate(
        headers,
        start=1,
    ):
        cell = worksheet.cell(
            row=header_row,
            column=column_index,
        )

        cell.value = header
        cell.font = Font(
            bold=True
        )
        cell.alignment = Alignment(
            horizontal="center"
        )

    # ---------------------------------------------------------
    # Table rows
    # ---------------------------------------------------------

    row_index = header_row + 1

    for item in report["items"]:

        values = [
            item["installment_number"],
            (
                item["due_date"].isoformat()
                if item["due_date"]
                else ""
            ),
            item["chit_name"],
            item["member_number"] or "",
            item["member_name"],
            item["customer_code"] or "",
            item["due_amount"],
            item["paid_amount"],
            item["remaining_amount"],
            item["status"].title(),
            item["days_overdue"],
        ]

        for column_index, value in enumerate(
            values,
            start=1,
        ):
            worksheet.cell(
                row=row_index,
                column=column_index,
                value=value,
            )

        row_index += 1

    # ---------------------------------------------------------
    # Column widths
    # ---------------------------------------------------------

    widths = {
        "A": 14,
        "B": 14,
        "C": 24,
        "D": 14,
        "E": 24,
        "F": 18,
        "G": 16,
        "H": 16,
        "I": 16,
        "J": 14,
        "K": 16,
    }

    for column, width in widths.items():
        worksheet.column_dimensions[
            column
        ].width = width

    # ---------------------------------------------------------
    # Freeze table header
    # ---------------------------------------------------------

    worksheet.freeze_panes = "A12"

    # ---------------------------------------------------------
    # Save workbook to memory
    # ---------------------------------------------------------

    output = BytesIO()

    workbook.save(output)

    output.seek(0)

    filename = (
        "due-report.xlsx"
    )

    return StreamingResponse(
        output,
        media_type=(
            "application/vnd.openxmlformats-officedocument"
            ".spreadsheetml.sheet"
        ),
        headers={
            "Content-Disposition": (
                f'attachment; filename="{filename}"'
            )
        },
    )


# ============================================================
# PDF Export
# ============================================================

@router.get(
    "/due/export/pdf"
)
def export_due_report_pdf(
    date_from: Optional[date] = Query(None),
    date_to: Optional[date] = Query(None),

    chit_id: Optional[int] = Query(
        None,
        gt=0,
    ),

    member_id: Optional[int] = Query(
        None,
        gt=0,
    ),

    report_status: str = Query(
        "all",
        alias="status",
        pattern="^(all|due|partial|overdue)$",
    ),

    search: Optional[str] = Query(None),

    db: Session = Depends(get_db),

    current_user: User = Depends(
        get_current_user
    ),
):
    if not current_user.company_id:
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail="User is not associated with a company.",
        )

    _validate_date_range(
        date_from,
        date_to,
    )

    try:
        report = get_due_report_export_data(
            db=db,
            company_id=current_user.company_id,

            date_from=date_from,
            date_to=date_to,

            chit_id=chit_id,
            member_id=member_id,

            report_status=report_status,

            search=search,
        )

    except ValueError as exc:
        _handle_company_error(exc)

    # ---------------------------------------------------------
    # PDF document
    # ---------------------------------------------------------

    output = BytesIO()

    document = SimpleDocTemplate(
        output,
        pagesize=landscape(A4),
        rightMargin=24,
        leftMargin=24,
        topMargin=24,
        bottomMargin=24,
    )

    styles = getSampleStyleSheet()

    title_style = styles["Title"]

    title_style.alignment = TA_LEFT

    normal_style = styles["Normal"]

    story = []

    # ---------------------------------------------------------
    # Company header
    # ---------------------------------------------------------

    story.append(
        Paragraph(
            report["company_name"],
            title_style,
        )
    )

    story.append(
        Paragraph(
            f"Company ID: {report['company_id']}",
            normal_style,
        )
    )

    story.append(
        Spacer(1, 8)
    )

    story.append(
        Paragraph(
            "DUE REPORT",
            styles["Heading2"],
        )
    )

    story.append(
        Spacer(1, 8)
    )

    # ---------------------------------------------------------
    # Summary
    # ---------------------------------------------------------

    summary = report["summary"]

    summary_data = [
        [
            "Total Due",
            "Total Paid",
            "Remaining",
            "Records",
            "Due",
            "Partial",
            "Overdue",
            "Members",
        ],
        [
            str(summary["total_due"]),
            str(summary["total_paid"]),
            str(summary["total_remaining"]),
            str(summary["total_records"]),
            str(summary["due_records"]),
            str(summary["partial_records"]),
            str(summary["overdue_records"]),
            str(summary["unique_members"]),
        ],
    ]

    summary_table = Table(
        summary_data,
        repeatRows=1,
    )

    summary_table.setStyle(
        TableStyle(
            [
                (
                    "BACKGROUND",
                    (0, 0),
                    (-1, 0),
                    colors.lightgrey,
                ),
                (
                    "FONTNAME",
                    (0, 0),
                    (-1, 0),
                    "Helvetica-Bold",
                ),
                (
                    "GRID",
                    (0, 0),
                    (-1, -1),
                    0.5,
                    colors.grey,
                ),
                (
                    "ALIGN",
                    (0, 0),
                    (-1, -1),
                    "CENTER",
                ),
                (
                    "VALIGN",
                    (0, 0),
                    (-1, -1),
                    "MIDDLE",
                ),
                (
                    "FONTSIZE",
                    (0, 0),
                    (-1, -1),
                    8,
                ),
            ]
        )
    )

    story.append(
        summary_table
    )

    story.append(
        Spacer(1, 12)
    )

    # ---------------------------------------------------------
    # Main table
    # ---------------------------------------------------------

    table_data = [
        [
            "Installment",
            "Due Date",
            "Chit",
            "Member No",
            "Member Name",
            "Customer Code",
            "Due",
            "Paid",
            "Remaining",
            "Status",
            "Days",
        ]
    ]

    for item in report["items"]:

        table_data.append(
            [
                str(
                    item[
                        "installment_number"
                    ]
                ),

                (
                    item["due_date"].isoformat()
                    if item["due_date"]
                    else ""
                ),

                Paragraph(
                    str(
                        item["chit_name"]
                    ),
                    normal_style,
                ),

                str(
                    item["member_number"]
                    or ""
                ),

                Paragraph(
                    str(
                        item["member_name"]
                    ),
                    normal_style,
                ),

                str(
                    item["customer_code"]
                    or ""
                ),

                str(
                    item["due_amount"]
                ),

                str(
                    item["paid_amount"]
                ),

                str(
                    item["remaining_amount"]
                ),

                str(
                    item["status"].title()
                ),

                str(
                    item["days_overdue"]
                ),
            ]
        )

    report_table = Table(
        table_data,
        repeatRows=1,
        colWidths=[
            52,
            68,
            95,
            65,
            105,
            75,
            62,
            62,
            70,
            58,
            48,
        ],
    )

    report_table.setStyle(
        TableStyle(
            [
                (
                    "BACKGROUND",
                    (0, 0),
                    (-1, 0),
                    colors.lightgrey,
                ),
                (
                    "FONTNAME",
                    (0, 0),
                    (-1, 0),
                    "Helvetica-Bold",
                ),
                (
                    "FONTSIZE",
                    (0, 0),
                    (-1, -1),
                    7,
                ),
                (
                    "GRID",
                    (0, 0),
                    (-1, -1),
                    0.35,
                    colors.grey,
                ),
                (
                    "VALIGN",
                    (0, 0),
                    (-1, -1),
                    "MIDDLE",
                ),
                (
                    "ALIGN",
                    (0, 0),
                    (-1, -1),
                    "CENTER",
                ),
                (
                    "ALIGN",
                    (4, 1),
                    (4, -1),
                    "LEFT",
                ),
                (
                    "ALIGN",
                    (2, 1),
                    (2, -1),
                    "LEFT",
                ),
            ]
        )
    )

    story.append(
        report_table
    )

    # ---------------------------------------------------------
    # Build PDF
    # ---------------------------------------------------------

    document.build(story)

    output.seek(0)

    filename = (
        "due-report.pdf"
    )

    return StreamingResponse(
        output,
        media_type="application/pdf",
        headers={
            "Content-Disposition": (
                f'attachment; filename="{filename}"'
            )
        },
    )