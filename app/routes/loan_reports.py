from datetime import date
from io import BytesIO
from typing import Optional

from fastapi import (
    APIRouter,
    Depends,
    HTTPException,
    Query,
    status,
)
from fastapi.responses import StreamingResponse

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import (
    SimpleDocTemplate,
    Table,
    TableStyle,
    Paragraph,
    Spacer,
)

from sqlalchemy.orm import Session

from app.core.database import get_db
from app.core.dependencies import get_current_user

from app.models.admin.user import User

from app.schemas.loan_report import (
    LoanReportResponse,
)

from app.services.loan_reports_service import (
    get_loan_report,
    get_loan_report_export_data,
)


from app.services.entitlement_service import require_company_feature

router = APIRouter(
    dependencies=[Depends(require_company_feature("loans"))],
    prefix="/reports",
    tags=["Reports"],
)


# =========================================================
# HELPERS
# =========================================================

def _validate_date_range(
    date_from: Optional[date],
    date_to: Optional[date],
):
    if (
        date_from
        and date_to
        and date_from > date_to
    ):
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=(
                "date_from cannot be later "
                "than date_to."
            ),
        )


def _handle_company_error(
    exc: ValueError,
):
    if str(exc) == "Company not found.":
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Company not found.",
        )

    raise exc


# =========================================================
# LOAN REPORT
# =========================================================

@router.get(
    "/loans",
    response_model=LoanReportResponse,
)
def loan_report(
    date_from: Optional[date] = Query(None),
    date_to: Optional[date] = Query(None),

    direction: str = Query(
        "all",
        pattern="^(all|given|taken)$",
    ),

    report_status: str = Query(
        "all",
        alias="status",
        pattern=(
            "^(all|active|partially_settled|"
            "closed|cancelled)$"
        ),
    ),

    search: Optional[str] = Query(None),

    as_of_date: Optional[date] = Query(None),

    page: int = Query(
        1,
        ge=1,
    ),

    page_size: int = Query(
        10,
        ge=1,
        le=100,
    ),

    db: Session = Depends(get_db),

    current_user: User = Depends(
        get_current_user
    ),
):
    # -----------------------------------------------------
    # COMPANY
    # -----------------------------------------------------

    if not current_user.company_id:
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail=(
                "User is not associated "
                "with a company."
            ),
        )

    # -----------------------------------------------------
    # DATE
    # -----------------------------------------------------

    _validate_date_range(
        date_from,
        date_to,
    )

    # -----------------------------------------------------
    # AS OF DATE VALIDATION
    # -----------------------------------------------------

    if (
        as_of_date
        and as_of_date > date.today()
    ):
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=(
                "as_of_date cannot be "
                "in the future."
            ),
        )

    # -----------------------------------------------------
    # REPORT
    # -----------------------------------------------------

    try:

        return get_loan_report(
            db=db,
            company_id=current_user.company_id,

            date_from=date_from,
            date_to=date_to,

            direction=direction,

            report_status=report_status,

            search=search,

            as_of_date=as_of_date,

            page=page,
            page_size=page_size,
        )

    except ValueError as exc:
        _handle_company_error(
            exc
        )


# =========================================================
# EXCEL EXPORT
# =========================================================

@router.get(
    "/loans/export/excel"
)
def export_loan_report_excel(
    date_from: Optional[date] = Query(None),
    date_to: Optional[date] = Query(None),

    direction: str = Query(
        "all",
        pattern="^(all|given|taken)$",
    ),

    report_status: str = Query(
        "all",
        alias="status",
        pattern=(
            "^(all|active|partially_settled|"
            "closed|cancelled)$"
        ),
    ),

    search: Optional[str] = Query(None),

    as_of_date: Optional[date] = Query(None),

    db: Session = Depends(get_db),

    current_user: User = Depends(
        get_current_user
    ),
):
    if not current_user.company_id:
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail=(
                "User is not associated "
                "with a company."
            ),
        )

    _validate_date_range(
        date_from,
        date_to,
    )

    if (
        as_of_date
        and as_of_date > date.today()
    ):
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=(
                "as_of_date cannot be "
                "in the future."
            ),
        )

    try:

        report = (
            get_loan_report_export_data(
                db=db,
                company_id=(
                    current_user.company_id
                ),

                date_from=date_from,
                date_to=date_to,

                direction=direction,

                report_status=report_status,

                search=search,

                as_of_date=as_of_date,
            )
        )

    except ValueError as exc:
        _handle_company_error(
            exc
        )

    # -----------------------------------------------------
    # WORKBOOK
    # -----------------------------------------------------

    workbook = Workbook()

    worksheet = workbook.active
    worksheet.title = "Loan Report"

    company_name = report[
        "company_name"
    ]

    company_id = report[
        "company_id"
    ]

    # -----------------------------------------------------
    # HEADER
    # -----------------------------------------------------

    worksheet["A1"] = company_name

    worksheet["A1"].font = Font(
        bold=True,
        size=16,
    )

    worksheet["A2"] = (
        f"Company ID: {company_id}"
    )

    worksheet["A4"] = "LOAN REPORT"

    worksheet["A4"].font = Font(
        bold=True,
        size=14,
    )

    worksheet["A5"] = (
        f"As of: {report['as_of_date']}"
    )

    # -----------------------------------------------------
    # SUMMARY
    # -----------------------------------------------------

    summary = report[
        "summary"
    ]

    worksheet["A7"] = "Total Loans"
    worksheet["B7"] = (
        summary["total_loans"]
    )

    worksheet["A8"] = "Given Loans"
    worksheet["B8"] = (
        summary["given_loans"]
    )

    worksheet["A9"] = "Taken Loans"
    worksheet["B9"] = (
        summary["taken_loans"]
    )

    worksheet["A10"] = "Total Principal"
    worksheet["B10"] = (
        summary["total_principal"]
    )

    worksheet["A11"] = "Total Disbursed"
    worksheet["B11"] = (
        summary["total_disbursed"]
    )

    worksheet["A12"] = "Principal Paid"
    worksheet["B12"] = (
        summary["principal_paid"]
    )

    worksheet["A13"] = "Interest Paid"
    worksheet["B13"] = (
        summary["interest_paid"]
    )

    worksheet["A14"] = "Discount"
    worksheet["B14"] = (
        summary["discount"]
    )

    worksheet["D7"] = "Principal Outstanding"
    worksheet["E7"] = (
        summary[
            "principal_outstanding"
        ]
    )

    worksheet["D8"] = "Interest Outstanding"
    worksheet["E8"] = (
        summary[
            "interest_outstanding"
        ]
    )

    worksheet["D9"] = "Total Outstanding"
    worksheet["E9"] = (
        summary[
            "total_outstanding"
        ]
    )

    worksheet["D10"] = "Cash Disbursed"
    worksheet["E10"] = (
        summary[
            "cash_disbursed"
        ]
    )

    worksheet["D11"] = "UPI Disbursed"
    worksheet["E11"] = (
        summary[
            "upi_disbursed"
        ]
    )

    worksheet["D12"] = "Cash Repaid"
    worksheet["E12"] = (
        summary[
            "cash_repaid"
        ]
    )

    worksheet["D13"] = "UPI Repaid"
    worksheet["E13"] = (
        summary[
            "upi_repaid"
        ]
    )

    worksheet["D14"] = "Unique Persons"
    worksheet["E14"] = (
        summary[
            "unique_persons"
        ]
    )

    # -----------------------------------------------------
    # TABLE
    # -----------------------------------------------------

    header_row = 16

    headers = [
        "ID",
        "Person Name",
        "Mobile",
        "Direction",
        "Principal",
        "Rate",
        "Interest Type",
        "Start Date",
        "Due Date",
        "Status",
        "Disbursed",
        "Principal Paid",
        "Interest Paid",
        "Discount",
        "Principal Outstanding",
        "Interest Outstanding",
        "Total Paid",
        "Total Outstanding",
        "Cash Disbursed",
        "UPI Disbursed",
        "Cash Repaid",
        "UPI Repaid",
        "Transactions",
        "Last Transaction",
    ]

    for column_index, header in enumerate(
        headers,
        start=1,
    ):

        cell = worksheet.cell(
            row=header_row,
            column=column_index,
        )

        cell.value = header

        cell.font = Font(
            bold=True
        )

        cell.alignment = Alignment(
            horizontal="center"
        )

    # -----------------------------------------------------
    # ROWS
    # -----------------------------------------------------

    row_index = (
        header_row + 1
    )

    for item in report[
        "items"
    ]:

        values = [
            item["id"],
            item["person_name"],
            item["person_mobile"] or "",
            item["direction"],
            item["principal_amount"],
            item["interest_rate"],
            item["interest_type"],
            (
                item["start_date"].isoformat()
                if item["start_date"]
                else ""
            ),
            (
                item["due_date"].isoformat()
                if item["due_date"]
                else ""
            ),
            item["status"],
            item["total_disbursed"],
            item["principal_paid"],
            item["interest_paid"],
            item["discount"],
            item["principal_outstanding"],
            item["interest_outstanding"],
            item["total_paid"],
            item["total_outstanding"],
            item["cash_disbursed"],
            item["upi_disbursed"],
            item["cash_repaid"],
            item["upi_repaid"],
            item["transaction_count"],
            (
                item[
                    "last_transaction_date"
                ].isoformat()
                if item[
                    "last_transaction_date"
                ]
                else ""
            ),
        ]

        for column_index, value in enumerate(
            values,
            start=1,
        ):

            worksheet.cell(
                row=row_index,
                column=column_index,
                value=value,
            )

        row_index += 1

    # -----------------------------------------------------
    # WIDTHS
    # -----------------------------------------------------

    widths = {
        "A": 10,
        "B": 24,
        "C": 16,
        "D": 12,
        "E": 16,
        "F": 10,
        "G": 16,
        "H": 14,
        "I": 14,
        "J": 18,
        "K": 16,
        "L": 16,
        "M": 16,
        "N": 14,
        "O": 20,
        "P": 20,
        "Q": 16,
        "R": 20,
        "S": 16,
        "T": 16,
        "U": 14,
        "V": 14,
        "W": 14,
        "X": 18,
    }

    for column, width in widths.items():

        worksheet.column_dimensions[
            column
        ].width = width

    worksheet.freeze_panes = (
        "A17"
    )

    # -----------------------------------------------------
    # OUTPUT
    # -----------------------------------------------------

    output = BytesIO()

    workbook.save(
        output
    )

    output.seek(0)

    return StreamingResponse(
        output,
        media_type=(
            "application/vnd.openxmlformats-officedocument"
            ".spreadsheetml.sheet"
        ),
        headers={
            "Content-Disposition": (
                'attachment; '
                'filename="loan-report.xlsx"'
            )
        },
    )


# =========================================================
# PDF EXPORT
# =========================================================

@router.get(
    "/loans/export/pdf"
)
def export_loan_report_pdf(
    date_from: Optional[date] = Query(None),
    date_to: Optional[date] = Query(None),

    direction: str = Query(
        "all",
        pattern="^(all|given|taken)$",
    ),

    report_status: str = Query(
        "all",
        alias="status",
        pattern=(
            "^(all|active|partially_settled|"
            "closed|cancelled)$"
        ),
    ),

    search: Optional[str] = Query(None),

    as_of_date: Optional[date] = Query(None),

    db: Session = Depends(get_db),

    current_user: User = Depends(
        get_current_user
    ),
):
    if not current_user.company_id:
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail=(
                "User is not associated "
                "with a company."
            ),
        )

    _validate_date_range(
        date_from,
        date_to,
    )

    if (
        as_of_date
        and as_of_date > date.today()
    ):
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=(
                "as_of_date cannot be "
                "in the future."
            ),
        )

    try:

        report = (
            get_loan_report_export_data(
                db=db,
                company_id=(
                    current_user.company_id
                ),

                date_from=date_from,
                date_to=date_to,

                direction=direction,

                report_status=report_status,

                search=search,

                as_of_date=as_of_date,
            )
        )

    except ValueError as exc:
        _handle_company_error(
            exc
        )

    # -----------------------------------------------------
    # PDF
    # -----------------------------------------------------

    output = BytesIO()

    document = SimpleDocTemplate(
        output,
        pagesize=landscape(A4),
        rightMargin=20,
        leftMargin=20,
        topMargin=20,
        bottomMargin=20,
    )

    styles = getSampleStyleSheet()

    story = []

    # -----------------------------------------------------
    # COMPANY
    # -----------------------------------------------------

    story.append(
        Paragraph(
            report["company_name"],
            styles["Title"],
        )
    )

    story.append(
        Paragraph(
            (
                f"Company ID: "
                f"{report['company_id']}"
            ),
            styles["Normal"],
        )
    )

    story.append(
        Spacer(1, 6)
    )

    story.append(
        Paragraph(
            "LOAN REPORT",
            styles["Heading2"],
        )
    )

    story.append(
        Paragraph(
            (
                f"As of: "
                f"{report['as_of_date']}"
            ),
            styles["Normal"],
        )
    )

    story.append(
        Spacer(1, 10)
    )

    # -----------------------------------------------------
    # SUMMARY
    # -----------------------------------------------------

    summary = report[
        "summary"
    ]

    summary_data = [
        [
            "Loans",
            "Given",
            "Taken",
            "Principal",
            "Disbursed",
            "Principal Paid",
            "Interest Paid",
            "Outstanding",
        ],
        [
            str(
                summary[
                    "total_loans"
                ]
            ),
            str(
                summary[
                    "given_loans"
                ]
            ),
            str(
                summary[
                    "taken_loans"
                ]
            ),
            str(
                summary[
                    "total_principal"
                ]
            ),
            str(
                summary[
                    "total_disbursed"
                ]
            ),
            str(
                summary[
                    "principal_paid"
                ]
            ),
            str(
                summary[
                    "interest_paid"
                ]
            ),
            str(
                summary[
                    "total_outstanding"
                ]
            ),
        ],
    ]

    summary_table = Table(
        summary_data,
        repeatRows=1,
    )

    summary_table.setStyle(
        TableStyle(
            [
                (
                    "BACKGROUND",
                    (0, 0),
                    (-1, 0),
                    colors.lightgrey,
                ),
                (
                    "FONTNAME",
                    (0, 0),
                    (-1, 0),
                    "Helvetica-Bold",
                ),
                (
                    "GRID",
                    (0, 0),
                    (-1, -1),
                    0.5,
                    colors.grey,
                ),
                (
                    "ALIGN",
                    (0, 0),
                    (-1, -1),
                    "CENTER",
                ),
                (
                    "FONTSIZE",
                    (0, 0),
                    (-1, -1),
                    7,
                ),
            ]
        )
    )

    story.append(
        summary_table
    )

    story.append(
        Spacer(1, 12)
    )

    # -----------------------------------------------------
    # REPORT TABLE
    # -----------------------------------------------------

    table_data = [
        [
            "ID",
            "Person",
            "Mobile",
            "Direction",
            "Principal",
            "Rate",
            "Type",
            "Start",
            "Due",
            "Status",
            "Disbursed",
            "Principal Paid",
            "Interest Paid",
            "Outstanding",
            "Cash Repaid",
            "UPI Repaid",
        ]
    ]

    normal_style = styles["Normal"]

    for item in report[
        "items"
    ]:

        table_data.append(
            [
                str(
                    item["id"]
                ),

                Paragraph(
                    str(
                        item[
                            "person_name"
                        ]
                    ),
                    normal_style,
                ),

                str(
                    item[
                        "person_mobile"
                    ]
                    or ""
                ),

                str(
                    item[
                        "direction"
                    ]
                ),

                str(
                    item[
                        "principal_amount"
                    ]
                ),

                str(
                    item[
                        "interest_rate"
                    ]
                ),

                str(
                    item[
                        "interest_type"
                    ]
                ),

                (
                    item[
                        "start_date"
                    ].isoformat()
                    if item[
                        "start_date"
                    ]
                    else ""
                ),

                (
                    item[
                        "due_date"
                    ].isoformat()
                    if item[
                        "due_date"
                    ]
                    else ""
                ),

                str(
                    item[
                        "status"
                    ]
                ),

                str(
                    item[
                        "total_disbursed"
                    ]
                ),

                str(
                    item[
                        "principal_paid"
                    ]
                ),

                str(
                    item[
                        "interest_paid"
                    ]
                ),

                str(
                    item[
                        "total_outstanding"
                    ]
                ),

                str(
                    item[
                        "cash_repaid"
                    ]
                ),

                str(
                    item[
                        "upi_repaid"
                    ]
                ),
            ]
        )

    report_table = Table(
        table_data,
        repeatRows=1,
        colWidths=[
            28,
            90,
            65,
            50,
            60,
            38,
            48,
            55,
            55,
            60,
            60,
            65,
            60,
            70,
            60,
            60,
        ],
    )

    report_table.setStyle(
        TableStyle(
            [
                (
                    "BACKGROUND",
                    (0, 0),
                    (-1, 0),
                    colors.lightgrey,
                ),
                (
                    "FONTNAME",
                    (0, 0),
                    (-1, 0),
                    "Helvetica-Bold",
                ),
                (
                    "FONTSIZE",
                    (0, 0),
                    (-1, -1),
                    6,
                ),
                (
                    "GRID",
                    (0, 0),
                    (-1, -1),
                    0.35,
                    colors.grey,
                ),
                (
                    "VALIGN",
                    (0, 0),
                    (-1, -1),
                    "MIDDLE",
                ),
                (
                    "ALIGN",
                    (0, 0),
                    (-1, -1),
                    "CENTER",
                ),
                (
                    "ALIGN",
                    (1, 1),
                    (1, -1),
                    "LEFT",
                ),
            ]
        )
    )

    story.append(
        report_table
    )

    # -----------------------------------------------------
    # BUILD
    # -----------------------------------------------------

    document.build(
        story
    )

    output.seek(0)

    return StreamingResponse(
        output,
        media_type="application/pdf",
        headers={
            "Content-Disposition": (
                'attachment; '
                'filename="loan-report.pdf"'
            )
        },
    )