from datetime import date
from io import BytesIO
from typing import Optional

from fastapi import (
    APIRouter,
    Depends,
    HTTPException,
    Query,
    status,
)

from fastapi.responses import StreamingResponse

from openpyxl import Workbook
from openpyxl.styles import Font

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import (
    SimpleDocTemplate,
    Table,
    TableStyle,
    Paragraph,
    Spacer,
)

from sqlalchemy.orm import Session

from app.core.database import get_db
from app.core.dependencies import get_current_user

from app.models.admin.user import User

from app.schemas.chits_reports import (
    ChitReportResponse,
)

from app.services.chits_reports_service import (
    get_chit_report,
    get_chit_report_export_data,
)


router = APIRouter(
    prefix="/reports",
    tags=["Reports"],
)


# =========================================================
# VALIDATION
# =========================================================

def _validate_date_range(
    date_from: Optional[date],
    date_to: Optional[date],
):

    if (
        date_from
        and date_to
        and date_from > date_to
    ):

        raise HTTPException(
            status_code=400,
            detail=(
                "date_from cannot be later "
                "than date_to."
            ),
        )


def _handle_company_error(
    exc: ValueError,
):

    if str(exc) == "Company not found.":

        raise HTTPException(
            status_code=404,
            detail="Company not found.",
        )

    raise exc


# =========================================================
# GET CHIT REPORT
# =========================================================

@router.get(
    "/chits",
    response_model=ChitReportResponse,
)
def chit_report(
    date_from: Optional[date] = Query(
        None
    ),

    date_to: Optional[date] = Query(
        None
    ),

    chit_id: Optional[int] = Query(
        None,
        gt=0,
    ),

    report_status: str = Query(
        "all",
        alias="status",
        pattern=(
            "^(all|draft|active|"
            "completed|cancelled)$"
        ),
    ),

    search: Optional[str] = Query(
        None
    ),

    page: int = Query(
        1,
        ge=1,
    ),

    page_size: int = Query(
        10,
        ge=1,
        le=100,
    ),

    db: Session = Depends(
        get_db
    ),

    current_user: User = Depends(
        get_current_user
    ),
):

    if not current_user.company_id:

        raise HTTPException(
            status_code=403,
            detail=(
                "User is not associated "
                "with a company."
            ),
        )

    _validate_date_range(
        date_from,
        date_to,
    )

    try:

        return get_chit_report(
            db=db,

            company_id=(
                current_user.company_id
            ),

            date_from=date_from,

            date_to=date_to,

            chit_id=chit_id,

            report_status=(
                report_status
            ),

            search=search,

            page=page,

            page_size=page_size,
        )

    except ValueError as exc:

        _handle_company_error(
            exc
        )


# =========================================================
# EXCEL EXPORT
# =========================================================

@router.get(
    "/chits/export/excel"
)
def export_chit_report_excel(
    date_from: Optional[date] = Query(
        None
    ),

    date_to: Optional[date] = Query(
        None
    ),

    chit_id: Optional[int] = Query(
        None,
        gt=0,
    ),

    report_status: str = Query(
        "all",
        alias="status",
        pattern=(
            "^(all|draft|active|"
            "completed|cancelled)$"
        ),
    ),

    search: Optional[str] = Query(
        None
    ),

    db: Session = Depends(
        get_db
    ),

    current_user: User = Depends(
        get_current_user
    ),
):

    if not current_user.company_id:

        raise HTTPException(
            status_code=403,
            detail=(
                "User is not associated "
                "with a company."
            ),
        )

    _validate_date_range(
        date_from,
        date_to,
    )

    try:

        report = (
            get_chit_report_export_data(
                db=db,

                company_id=(
                    current_user.company_id
                ),

                date_from=date_from,

                date_to=date_to,

                chit_id=chit_id,

                report_status=(
                    report_status
                ),

                search=search,
            )
        )

    except ValueError as exc:

        _handle_company_error(
            exc
        )

    # =====================================================
    # WORKBOOK
    # =====================================================

    workbook = Workbook()

    sheet = workbook.active

    sheet.title = "Chit Report"

    # =====================================================
    # HEADER
    # =====================================================

    sheet["A1"] = (
        report["company_name"]
    )

    sheet["A1"].font = Font(
        bold=True,
        size=16,
    )

    sheet["A2"] = (
        f"Company ID: "
        f"{report['company_id']}"
    )

    sheet["A4"] = (
        "CHIT REPORT"
    )

    sheet["A4"].font = Font(
        bold=True,
        size=14,
    )

    # =====================================================
    # SUMMARY
    # =====================================================

    summary = report[
        "summary"
    ]

    summary_rows = [
        (
            "Total Chits",
            summary[
                "total_chits"
            ],
        ),

        (
            "Draft Chits",
            summary[
                "draft_chits"
            ],
        ),

        (
            "Active Chits",
            summary[
                "active_chits"
            ],
        ),

        (
            "Completed Chits",
            summary[
                "completed_chits"
            ],
        ),

        (
            "Total Chit Value",
            summary[
                "total_chit_value"
            ],
        ),

        (
            "Total Members",
            summary[
                "total_members"
            ],
        ),

        (
            "Scheduled Amount",
            summary[
                "total_scheduled_amount"
            ],
        ),

        (
            "Collected Amount",
            summary[
                "total_collected_amount"
            ],
        ),

        (
            "Outstanding Amount",
            summary[
                "total_outstanding_amount"
            ],
        ),

        (
            "Collection Rate",
            summary[
                "collection_rate"
            ],
        ),

        (
            "Total Lifts",
            summary[
                "total_lifts"
            ],
        ),

        (
            "Completed Payouts",
            summary[
                "completed_payouts"
            ],
        ),

        (
            "Pending Payouts",
            summary[
                "pending_payouts"
            ],
        ),

        (
            "Total Lifted Amount",
            summary[
                "total_lifted_amount"
            ],
        ),

        (
            "Total Discount",
            summary[
                "total_discount_amount"
            ],
        ),
    ]

    row = 6

    for label, value in summary_rows:

        sheet.cell(
            row=row,
            column=1,
            value=label,
        )

        sheet.cell(
            row=row,
            column=2,
            value=value,
        )

        row += 1

    # =====================================================
    # CHIT DETAILS
    # =====================================================

    row += 1

    sheet.cell(
        row=row,
        column=1,
        value="CHIT DETAILS",
    ).font = Font(
        bold=True
    )

    row += 1

    headers = [
        "ID",
        "Name",
        "Code",
        "Chit Value",
        "Monthly EMI",
        "Duration",
        "Max Members",
        "Start Date",
        "Status",
        "Members",
        "Confirmed",
        "Active",
        "Scheduled",
        "Collected",
        "Outstanding",
        "Collection %",
        "Lifts",
        "Lifted Amount",
        "Discount",
        "Payout Completed",
        "Payout Pending",
        "Last Collection",
        "Last Lift",
    ]

    for index, header in enumerate(
        headers,
        start=1,
    ):

        cell = sheet.cell(
            row=row,
            column=index,
            value=header,
        )

        cell.font = Font(
            bold=True
        )

    row += 1

    for item in report[
        "items"
    ]:

        values = [
            item["id"],
            item["name"],
            item["code"],
            item["chit_value"],
            item["monthly_emi"],
            item["duration_months"],
            item["max_members"],
            (
                item["start_date"].isoformat()
                if item["start_date"]
                else ""
            ),
            item["status"],
            item["member_count"],
            item["confirmed_members"],
            item["active_members"],
            item["scheduled_amount"],
            item["collected_amount"],
            item["outstanding_amount"],
            item["collection_rate"],
            item["total_lifts"],
            item["lifted_amount"],
            item["discount_amount"],
            item["payout_completed"],
            item["payout_pending"],
            (
                item[
                    "last_collection_date"
                ].isoformat()
                if item[
                    "last_collection_date"
                ]
                else ""
            ),
            (
                item[
                    "last_lift_date"
                ].isoformat()
                if item[
                    "last_lift_date"
                ]
                else ""
            ),
        ]

        for index, value in enumerate(
            values,
            start=1,
        ):

            sheet.cell(
                row=row,
                column=index,
                value=value,
            )

        row += 1

    # =====================================================
    # MONTHLY DETAILS
    # =====================================================

    row += 2

    sheet.cell(
        row=row,
        column=1,
        value="MONTHLY CHIT DETAILS",
    ).font = Font(
        bold=True
    )

    row += 1

    monthly_headers = [
        "Chit",
        "Code",
        "Installment",
        "Month Date",
        "EMI",
        "Planned Lift",
        "Scheduled",
        "Collected",
        "Outstanding",
        "Collection %",
        "Paid",
        "Partial",
        "Pending",
        "Lift Status",
        "Lifted",
        "Discount",
    ]

    for index, header in enumerate(
        monthly_headers,
        start=1,
    ):

        sheet.cell(
            row=row,
            column=index,
            value=header,
        ).font = Font(
            bold=True
        )

    row += 1

    for item in report[
        "monthly_summary"
    ]:

        values = [
            item["chit_name"],
            item["chit_code"],
            item["installment_number"],
            (
                item["month_date"].isoformat()
                if item["month_date"]
                else ""
            ),
            item["monthly_emi"],
            item["monthly_lift_amount"],
            item["scheduled_amount"],
            item["collected_amount"],
            item["outstanding_amount"],
            item["collection_rate"],
            item["paid_installments"],
            item["partial_installments"],
            item["pending_installments"],
            item["lift_status"],
            item["lifted_amount"],
            item["discount_amount"],
        ]

        for index, value in enumerate(
            values,
            start=1,
        ):

            sheet.cell(
                row=row,
                column=index,
                value=value,
            )

        row += 1

    # =====================================================
    # WIDTHS
    # =====================================================

    for column in sheet.columns:

        max_length = 0

        column_letter = (
            column[0].column_letter
        )

        for cell in column:

            value = (
                str(cell.value)
                if cell.value is not None
                else ""
            )

            max_length = max(
                max_length,
                len(value),
            )

        sheet.column_dimensions[
            column_letter
        ].width = min(
            max_length + 2,
            35,
        )

    sheet.freeze_panes = "A1"

    # =====================================================
    # OUTPUT
    # =====================================================

    output = BytesIO()

    workbook.save(
        output
    )

    output.seek(0)

    return StreamingResponse(
        output,

        media_type=(
            "application/vnd.openxmlformats-officedocument"
            ".spreadsheetml.sheet"
        ),

        headers={
            "Content-Disposition": (
                'attachment; '
                'filename="chit-report.xlsx"'
            )
        },
    )


# =========================================================
# PDF EXPORT
# =========================================================

@router.get(
    "/chits/export/pdf"
)
def export_chit_report_pdf(
    date_from: Optional[date] = Query(
        None
    ),

    date_to: Optional[date] = Query(
        None
    ),

    chit_id: Optional[int] = Query(
        None,
        gt=0,
    ),

    report_status: str = Query(
        "all",
        alias="status",
        pattern=(
            "^(all|draft|active|"
            "completed|cancelled)$"
        ),
    ),

    search: Optional[str] = Query(
        None
    ),

    db: Session = Depends(
        get_db
    ),

    current_user: User = Depends(
        get_current_user
    ),
):

    if not current_user.company_id:

        raise HTTPException(
            status_code=403,
            detail=(
                "User is not associated "
                "with a company."
            ),
        )

    _validate_date_range(
        date_from,
        date_to,
    )

    try:

        report = (
            get_chit_report_export_data(
                db=db,

                company_id=(
                    current_user.company_id
                ),

                date_from=date_from,

                date_to=date_to,

                chit_id=chit_id,

                report_status=(
                    report_status
                ),

                search=search,
            )
        )

    except ValueError as exc:

        _handle_company_error(
            exc
        )

    # =====================================================
    # PDF
    # =====================================================

    output = BytesIO()

    document = SimpleDocTemplate(
        output,

        pagesize=landscape(A4),

        rightMargin=20,
        leftMargin=20,

        topMargin=20,
        bottomMargin=20,
    )

    styles = (
        getSampleStyleSheet()
    )

    story = []

    # =====================================================
    # HEADER
    # =====================================================

    story.append(
        Paragraph(
            report[
                "company_name"
            ],
            styles["Title"],
        )
    )

    story.append(
        Paragraph(
            (
                f"Company ID: "
                f"{report['company_id']}"
            ),
            styles["Normal"],
        )
    )

    story.append(
        Spacer(1, 6)
    )

    story.append(
        Paragraph(
            "CHIT REPORT",
            styles["Heading2"],
        )
    )

    story.append(
        Spacer(1, 10)
    )

    # =====================================================
    # SUMMARY
    # =====================================================

    summary = report[
        "summary"
    ]

    summary_data = [
        [
            "Chits",
            "Chit Value",
            "Members",
            "Scheduled",
            "Collected",
            "Outstanding",
            "Collection %",
            "Lifts",
            "Lifted",
            "Discount",
        ],

        [
            str(
                summary[
                    "total_chits"
                ]
            ),

            str(
                summary[
                    "total_chit_value"
                ]
            ),

            str(
                summary[
                    "total_members"
                ]
            ),

            str(
                summary[
                    "total_scheduled_amount"
                ]
            ),

            str(
                summary[
                    "total_collected_amount"
                ]
            ),

            str(
                summary[
                    "total_outstanding_amount"
                ]
            ),

            str(
                summary[
                    "collection_rate"
                ]
            ),

            str(
                summary[
                    "total_lifts"
                ]
            ),

            str(
                summary[
                    "total_lifted_amount"
                ]
            ),

            str(
                summary[
                    "total_discount_amount"
                ]
            ),
        ],
    ]

    summary_table = Table(
        summary_data,
        repeatRows=1,
    )

    summary_table.setStyle(
        TableStyle(
            [
                (
                    "BACKGROUND",
                    (0, 0),
                    (-1, 0),
                    colors.lightgrey,
                ),

                (
                    "FONTNAME",
                    (0, 0),
                    (-1, 0),
                    "Helvetica-Bold",
                ),

                (
                    "GRID",
                    (0, 0),
                    (-1, -1),
                    0.4,
                    colors.grey,
                ),

                (
                    "ALIGN",
                    (0, 0),
                    (-1, -1),
                    "CENTER",
                ),

                (
                    "FONTSIZE",
                    (0, 0),
                    (-1, -1),
                    7,
                ),
            ]
        )
    )

    story.append(
        summary_table
    )

    story.append(
        Spacer(1, 14)
    )

    # =====================================================
    # CHIT TABLE
    # =====================================================

    story.append(
        Paragraph(
            "Chit Details",
            styles["Heading3"],
        )
    )

    chit_data = [
        [
            "Name",
            "Code",
            "Value",
            "Members",
            "Status",
            "Scheduled",
            "Collected",
            "Outstanding",
            "Collection %",
            "Lifts",
            "Lifted",
            "Discount",
        ]
    ]

    for item in report[
        "items"
    ]:

        chit_data.append(
            [
                str(item["name"]),
                str(item["code"]),
                str(item["chit_value"]),
                str(item["member_count"]),
                str(item["status"]),
                str(item["scheduled_amount"]),
                str(item["collected_amount"]),
                str(item["outstanding_amount"]),
                str(item["collection_rate"]),
                str(item["total_lifts"]),
                str(item["lifted_amount"]),
                str(item["discount_amount"]),
            ]
        )

    chit_table = Table(
        chit_data,
        repeatRows=1,
    )

    chit_table.setStyle(
        TableStyle(
            [
                (
                    "BACKGROUND",
                    (0, 0),
                    (-1, 0),
                    colors.lightgrey,
                ),

                (
                    "FONTNAME",
                    (0, 0),
                    (-1, 0),
                    "Helvetica-Bold",
                ),

                (
                    "GRID",
                    (0, 0),
                    (-1, -1),
                    0.35,
                    colors.grey,
                ),

                (
                    "FONTSIZE",
                    (0, 0),
                    (-1, -1),
                    6.5,
                ),

                (
                    "ALIGN",
                    (2, 1),
                    (-1, -1),
                    "CENTER",
                ),
            ]
        )
    )

    story.append(
        chit_table
    )

    story.append(
        Spacer(1, 14)
    )

    # =====================================================
    # MONTHLY TABLE
    # =====================================================

    story.append(
        Paragraph(
            "Monthly Chit Details",
            styles["Heading3"],
        )
    )

    monthly_data = [
        [
            "Chit",
            "Month",
            "Date",
            "EMI",
            "Scheduled",
            "Collected",
            "Outstanding",
            "Collection %",
            "Paid",
            "Partial",
            "Pending",
            "Lift",
            "Lifted",
            "Discount",
        ]
    ]

    for item in report[
        "monthly_summary"
    ]:

        monthly_data.append(
            [
                str(
                    item[
                        "chit_name"
                    ]
                ),

                str(
                    item[
                        "installment_number"
                    ]
                ),

                (
                    item[
                        "month_date"
                    ].isoformat()
                    if item[
                        "month_date"
                    ]
                    else ""
                ),

                str(
                    item[
                        "monthly_emi"
                    ]
                ),

                str(
                    item[
                        "scheduled_amount"
                    ]
                ),

                str(
                    item[
                        "collected_amount"
                    ]
                ),

                str(
                    item[
                        "outstanding_amount"
                    ]
                ),

                str(
                    item[
                        "collection_rate"
                    ]
                ),

                str(
                    item[
                        "paid_installments"
                    ]
                ),

                str(
                    item[
                        "partial_installments"
                    ]
                ),

                str(
                    item[
                        "pending_installments"
                    ]
                ),

                str(
                    item[
                        "lift_status"
                    ]
                ),

                str(
                    item[
                        "lifted_amount"
                    ]
                ),

                str(
                    item[
                        "discount_amount"
                    ]
                ),
            ]
        )

    monthly_table = Table(
        monthly_data,
        repeatRows=1,
    )

    monthly_table.setStyle(
        TableStyle(
            [
                (
                    "BACKGROUND",
                    (0, 0),
                    (-1, 0),
                    colors.lightgrey,
                ),

                (
                    "FONTNAME",
                    (0, 0),
                    (-1, 0),
                    "Helvetica-Bold",
                ),

                (
                    "GRID",
                    (0, 0),
                    (-1, -1),
                    0.35,
                    colors.grey,
                ),

                (
                    "FONTSIZE",
                    (0, 0),
                    (-1, -1),
                    6,
                ),

                (
                    "ALIGN",
                    (1, 1),
                    (-1, -1),
                    "CENTER",
                ),
            ]
        )
    )

    story.append(
        monthly_table
    )

    document.build(
        story
    )

    output.seek(0)

    return StreamingResponse(
        output,

        media_type="application/pdf",

        headers={
            "Content-Disposition": (
                'attachment; '
                'filename="chit-report.pdf"'
            )
        },
    )