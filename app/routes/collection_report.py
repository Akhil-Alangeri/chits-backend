from datetime import date
from typing import Optional
from fastapi.responses import StreamingResponse
from io import BytesIO

from reportlab.lib import colors
from reportlab.lib.pagesizes import landscape, A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import (
    SimpleDocTemplate,
    Table,
    TableStyle,
    Paragraph,
    Spacer,
)

from openpyxl import Workbook

from fastapi import (
    APIRouter,
    Depends,
    HTTPException,
    Query,
    status,
)
from sqlalchemy.orm import Session

from app.core.database import get_db
from app.core.dependencies import get_current_user

from app.models.admin.user import User

from app.schemas.collection_reports import (
    CollectionReportResponse,
)

from app.services.collection_reports_service import (
    get_collection_report,
    get_collection_report_export_data
)


router = APIRouter(
    prefix="/reports",
    tags=["Reports"],
)


# =========================================================
# COLLECTION REPORT
# =========================================================

@router.get(
    "/collections",
    response_model=CollectionReportResponse,
)
def collection_report(

    date_from: Optional[date] = Query(
        None
    ),

    date_to: Optional[date] = Query(
        None
    ),

    chit_id: Optional[int] = Query(
        None,
        gt=0,
    ),

    member_id: Optional[int] = Query(
        None,
        gt=0,
    ),

    payment_method: Optional[str] = Query(
        None
    ),

    report_status: str = Query(
        "all",
        alias="status",
        pattern="^(all|paid|partial|due)$",
    ),

    search: Optional[str] = Query(
        None
    ),

    page: int = Query(
        1,
        ge=1,
    ),

    page_size: int = Query(
        10,
        ge=1,
        le=100,
    ),

    db: Session = Depends(
        get_db
    ),

    current_user: User = Depends(
        get_current_user
    ),
):

    # =====================================================
    # VERIFY COMPANY
    # =====================================================

    if not current_user.company_id:

        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail="User is not associated with a company.",
        )

    # =====================================================
    # DATE VALIDATION
    # =====================================================

    if (
        date_from
        and date_to
        and date_from > date_to
    ):

        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="date_from cannot be later than date_to.",
        )

    # =====================================================
    # GET REPORT
    # =====================================================
    #
    # IMPORTANT:
    #
    # company_id comes ONLY from current_user.company_id.
    #
    # Do NOT add company_id as a query parameter.
    #
    # =====================================================

    try:

        return get_collection_report(

            db=db,

            company_id=
                current_user.company_id,

            date_from=date_from,

            date_to=date_to,

            chit_id=chit_id,

            member_id=member_id,

            payment_method=payment_method,

            report_status=report_status,

            search=search,

            page=page,

            page_size=page_size,
        )

    except ValueError as exc:

        if str(exc) == "Company not found.":

            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail="Company not found.",
            )

        raise
    
    
    
    
# ======================================================
# SHEEETS FORMAT
# ======================================================

@router.get(
    "/collections/export/excel"
)

def export_collection_report_excel(

    date_from: Optional[date] = Query(None),

    date_to: Optional[date] = Query(None),

    chit_id: Optional[int] = Query(
        None,
        gt=0,
    ),

    member_id: Optional[int] = Query(
        None,
        gt=0,
    ),

    payment_method: Optional[str] = Query(None),

    report_status: str = Query(
        "all",
        alias="status",
        pattern="^(all|paid|partial|due)$",
    ),

    search: Optional[str] = Query(None),

    db: Session = Depends(get_db),

    current_user: User = Depends(get_current_user),
):

    if not current_user.company_id:

        raise HTTPException(
            status_code=403,
            detail="User is not associated with a company.",
        )

    report = get_collection_report_export_data(
        db=db,
        company_id=current_user.company_id,
        date_from=date_from,
        date_to=date_to,
        chit_id=chit_id,
        member_id=member_id,
        payment_method=payment_method,
        report_status=report_status,
        search=search,
    )

    output = BytesIO()

    # Excel generation here
    # using openpyxl


    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Collection Report"

    # Company details
    sheet["A1"] = report["company_name"]
    sheet["A2"] = f"Company ID: {report['company_id']}"
    sheet["A3"] = "Collection Report"

    headers = [
        "Date",
        "Receipt No",
        "Member",
        "Member Number",
        "Chit",
        "Installment",
        "Due Amount",
        "Paid Amount",
        "Remaining",
        "Payment Method",
        "Reference",
        "Collected By",
        "Status",
    ]

    header_row = 5

    for column, header in enumerate(
        headers,
        start=1,
    ):
        sheet.cell(
            row=header_row,
            column=column,
            value=header,
        )

    for row_index, item in enumerate(
        report["items"],
        start=header_row + 1,
    ):

        values = [
            item["date"],
            item["receipt_no"],
            item["member_name"],
            item["member_number"],
            item["chit_name"],
            item["installment_number"],
            item["due_amount"],
            item["paid_amount"],
            item["remaining_amount"],
            item["payment_method"],
            item["reference_number"],
            item["collected_by"],
            item["status"],
        ]

        for column, value in enumerate(
            values,
            start=1,
        ):

            sheet.cell(
                row=row_index,
                column=column,
                value=value,
            )

    # Auto width
    for column_cells in sheet.columns:

        max_length = 0

        column_letter = column_cells[0].column_letter

        for cell in column_cells:

            if cell.value is not None:

                max_length = max(
                    max_length,
                    len(str(cell.value)),
                )

        sheet.column_dimensions[
            column_letter
        ].width = min(
            max_length + 2,
            35,
        )

    workbook.save(output)

    output.seek(0)

    return StreamingResponse(
        output,
        media_type=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
        headers={
            "Content-Disposition":
                'attachment; filename="collection-report.xlsx"'
        },
    )
    


# ====================================================================
# PDF FORMAT
# ====================================================================

@router.get(
    "/collections/export/pdf"
)
def export_collection_report_pdf(
    date_from: Optional[date] = Query(None),
    date_to: Optional[date] = Query(None),

    chit_id: Optional[int] = Query(
        None,
        gt=0,
    ),

    member_id: Optional[int] = Query(
        None,
        gt=0,
    ),

    payment_method: Optional[str] = Query(None),

    report_status: str = Query(
        "all",
        alias="status",
        pattern="^(all|paid|partial|due)$",
    ),

    search: Optional[str] = Query(None),

    db: Session = Depends(get_db),

    current_user: User = Depends(get_current_user),
):

    # =====================================================
    # VERIFY COMPANY
    # =====================================================

    if not current_user.company_id:
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail="User is not associated with a company.",
        )

    # =====================================================
    # DATE VALIDATION
    # =====================================================

    if date_from and date_to and date_from > date_to:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="date_from cannot be later than date_to.",
        )

    # =====================================================
    # GET REPORT DATA
    # =====================================================

    try:
        report = get_collection_report_export_data(
            db=db,
            company_id=current_user.company_id,
            date_from=date_from,
            date_to=date_to,
            chit_id=chit_id,
            member_id=member_id,
            payment_method=payment_method,
            report_status=report_status,
            search=search,
        )

    except ValueError as exc:

        if str(exc) == "Company not found.":
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail="Company not found.",
            )

        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=str(exc),
        )

    # =====================================================
    # CREATE PDF
    # =====================================================

    output = BytesIO()

    document = SimpleDocTemplate(
        output,
        pagesize=landscape(A4),

        leftMargin=20,
        rightMargin=20,
        topMargin=20,
        bottomMargin=20,

        title="Collection Report",
        author=report["company_name"],
    )

    styles = getSampleStyleSheet()

    # =====================================================
    # CUSTOM STYLES
    # =====================================================

    company_style = styles["Title"].clone(
        "CollectionReportCompany"
    )

    company_style.fontName = "Helvetica-Bold"
    company_style.fontSize = 16
    company_style.leading = 20
    company_style.spaceAfter = 3

    company_id_style = styles["Normal"].clone(
        "CollectionReportCompanyId"
    )

    company_id_style.fontName = "Helvetica"
    company_id_style.fontSize = 8
    company_id_style.leading = 10

    report_title_style = styles["Heading2"].clone(
        "CollectionReportTitle"
    )

    report_title_style.fontName = "Helvetica-Bold"
    report_title_style.fontSize = 13
    report_title_style.leading = 16
    report_title_style.spaceBefore = 8
    report_title_style.spaceAfter = 8

    cell_style = styles["Normal"].clone(
        "CollectionReportCell"
    )

    cell_style.fontName = "Helvetica"
    cell_style.fontSize = 6.5
    cell_style.leading = 8

    header_style = styles["Normal"].clone(
        "CollectionReportHeader"
    )

    header_style.fontName = "Helvetica-Bold"
    header_style.fontSize = 6.5
    header_style.leading = 8
    header_style.textColor = colors.white

    # =====================================================
    # PDF ELEMENTS
    # =====================================================

    elements = []

    # =====================================================
    # COMPANY HEADER
    # =====================================================

    company_name = report.get("company_name") or "Company"

    company_id = report.get("company_id") or ""

    elements.append(
        Paragraph(
            company_name,
            company_style,
        )
    )

    elements.append(
        Paragraph(
            f"Company ID: {company_id}",
            company_id_style,
        )
    )

    elements.append(
        Spacer(1, 8)
    )

    elements.append(
        Paragraph(
            "COLLECTION REPORT",
            report_title_style,
        )
    )

    # =====================================================
    # REPORT PERIOD
    # =====================================================

    period_text = ""

    if date_from and date_to:
        period_text = (
            f"Period: {date_from.strftime('%d-%m-%Y')} "
            f"to {date_to.strftime('%d-%m-%Y')}"
        )

    elif date_from:
        period_text = (
            f"From: {date_from.strftime('%d-%m-%Y')}"
        )

    elif date_to:
        period_text = (
            f"Up to: {date_to.strftime('%d-%m-%Y')}"
        )

    if period_text:
        elements.append(
            Paragraph(
                period_text,
                company_id_style,
            )
        )

        elements.append(
            Spacer(1, 8)
        )

    # =====================================================
    # SUMMARY
    # =====================================================

    summary = report.get("summary", {})

    summary_data = [
        [
            Paragraph(
                "<b>Total Collected</b>",
                cell_style,
            ),
            Paragraph(
                "<b>Total Payments</b>",
                cell_style,
            ),
            Paragraph(
                "<b>Unique Members</b>",
                cell_style,
            ),
            Paragraph(
                "<b>Pending Amount</b>",
                cell_style,
            ),
        ],
        [
            Paragraph(
                f"₹ {summary.get('total_collected', 0):,}",
                cell_style,
            ),
            Paragraph(
                str(summary.get("total_payments", 0)),
                cell_style,
            ),
            Paragraph(
                str(summary.get("unique_members", 0)),
                cell_style,
            ),
            Paragraph(
                f"₹ {summary.get('pending_amount', 0):,}",
                cell_style,
            ),
        ],
    ]

    summary_table = Table(
        summary_data,
        colWidths=[
            180,
            150,
            150,
            180,
        ],
        repeatRows=1,
    )

    summary_table.setStyle(
        TableStyle(
            [
                (
                    "BACKGROUND",
                    (0, 0),
                    (-1, 0),
                    colors.HexColor("#e2e8f0"),
                ),
                (
                    "GRID",
                    (0, 0),
                    (-1, -1),
                    0.4,
                    colors.HexColor("#cbd5e1"),
                ),
                (
                    "VALIGN",
                    (0, 0),
                    (-1, -1),
                    "MIDDLE",
                ),
                (
                    "ALIGN",
                    (0, 0),
                    (-1, -1),
                    "CENTER",
                ),
                (
                    "TOPPADDING",
                    (0, 0),
                    (-1, -1),
                    5,
                ),
                (
                    "BOTTOMPADDING",
                    (0, 0),
                    (-1, -1),
                    5,
                ),
            ]
        )
    )

    elements.append(summary_table)

    elements.append(
        Spacer(1, 12)
    )

    # =====================================================
    # TABLE HEADER
    # =====================================================

    data = [
        [
            Paragraph("Date", header_style),
            Paragraph("Receipt", header_style),
            Paragraph("Member", header_style),
            Paragraph("Chit", header_style),
            Paragraph("Installment", header_style),
            Paragraph("Due", header_style),
            Paragraph("Paid", header_style),
            Paragraph("Remaining", header_style),
            Paragraph("Method", header_style),
            Paragraph("Reference", header_style),
            Paragraph("Status", header_style),
        ]
    ]

    # =====================================================
    # TABLE DATA
    # =====================================================

    for item in report.get("items", []):

        date_value = item.get("date")

        if date_value:
            date_text = str(date_value)
        else:
            date_text = ""

        installment_number = (
            item.get("installment_number")
            if item.get("installment_number") is not None
            else ""
        )

        due_amount = item.get("due_amount") or 0

        paid_amount = item.get("paid_amount") or 0

        remaining_amount = (
            item.get("remaining_amount") or 0
        )

        data.append(
            [
                Paragraph(
                    date_text,
                    cell_style,
                ),

                Paragraph(
                    str(item.get("receipt_no") or ""),
                    cell_style,
                ),

                Paragraph(
                    str(item.get("member_name") or ""),
                    cell_style,
                ),

                Paragraph(
                    str(item.get("chit_name") or ""),
                    cell_style,
                ),

                Paragraph(
                    str(installment_number),
                    cell_style,
                ),

                Paragraph(
                    f"₹ {due_amount:,}",
                    cell_style,
                ),

                Paragraph(
                    f"₹ {paid_amount:,}",
                    cell_style,
                ),

                Paragraph(
                    f"₹ {remaining_amount:,}",
                    cell_style,
                ),

                Paragraph(
                    str(item.get("payment_method") or ""),
                    cell_style,
                ),

                Paragraph(
                    str(item.get("reference_number") or ""),
                    cell_style,
                ),

                Paragraph(
                    str(item.get("status") or "").title(),
                    cell_style,
                ),
            ]
        )

    # =====================================================
    # REPORT TABLE
    # =====================================================

    table = Table(
        data,

        # Total = 785 points approximately.
        # A4 landscape usable width is around 802 points
        # with 20 point margins.
        colWidths=[
            55,   # Date
            60,   # Receipt
            100,  # Member
            95,   # Chit
            60,   # Installment
            65,   # Due
            65,   # Paid
            70,   # Remaining
            65,   # Method
            95,   # Reference
            55,   # Status
        ],

        repeatRows=1,

        splitByRow=1,
    )

    # =====================================================
    # TABLE STYLE
    # =====================================================

    table.setStyle(
        TableStyle(
            [
                (
                    "BACKGROUND",
                    (0, 0),
                    (-1, 0),
                    colors.HexColor("#1e293b"),
                ),

                (
                    "TEXTCOLOR",
                    (0, 0),
                    (-1, 0),
                    colors.white,
                ),

                (
                    "FONTNAME",
                    (0, 0),
                    (-1, 0),
                    "Helvetica-Bold",
                ),

                (
                    "FONTSIZE",
                    (0, 0),
                    (-1, -1),
                    6.5,
                ),

                (
                    "GRID",
                    (0, 0),
                    (-1, -1),
                    0.35,
                    colors.HexColor("#cbd5e1"),
                ),

                (
                    "VALIGN",
                    (0, 0),
                    (-1, -1),
                    "MIDDLE",
                ),

                (
                    "ROWBACKGROUNDS",
                    (0, 1),
                    (-1, -1),
                    [
                        colors.white,
                        colors.HexColor("#f8fafc"),
                    ],
                ),

                (
                    "LEFTPADDING",
                    (0, 0),
                    (-1, -1),
                    4,
                ),

                (
                    "RIGHTPADDING",
                    (0, 0),
                    (-1, -1),
                    4,
                ),

                (
                    "TOPPADDING",
                    (0, 0),
                    (-1, -1),
                    4,
                ),

                (
                    "BOTTOMPADDING",
                    (0, 0),
                    (-1, -1),
                    4,
                ),
            ]
        )
    )

    elements.append(table)

    # =====================================================
    # BUILD PDF
    # =====================================================

    document.build(elements)

    output.seek(0)

    # =====================================================
    # RETURN PDF
    # =====================================================

    return StreamingResponse(
        output,
        media_type="application/pdf",
        headers={
            "Content-Disposition": (
                'attachment; filename="collection-report.pdf"'
            )
        },
    )
    
    
