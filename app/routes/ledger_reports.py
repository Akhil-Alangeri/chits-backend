from datetime import date
from io import BytesIO
from typing import Optional

from fastapi import (
    APIRouter,
    Depends,
    HTTPException,
    Query,
)

from fastapi.responses import StreamingResponse

from openpyxl import Workbook
from openpyxl.styles import (
    Font,
    Alignment,
)

from reportlab.lib import colors
from reportlab.lib.pagesizes import (
    A4,
    landscape,
)

from reportlab.lib.styles import (
    getSampleStyleSheet,
)

from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
)

from sqlalchemy.orm import Session

from app.core.database import get_db
from app.core.dependencies import get_current_user

from app.models.admin.user import User

from app.schemas.ledger_report import (
    LedgerReportResponse,
)

from app.services.ledger_report_service import (
    get_ledger_report,
    get_ledger_report_export_data,
)


# =========================================================
# ROUTER
# =========================================================

router = APIRouter(
    prefix="/reports",
    tags=["Reports"],
)


# =========================================================
# DATE VALIDATION
# =========================================================

def _validate_dates(
    date_from,
    date_to,
):

    if (
        date_from
        and date_to
        and date_from > date_to
    ):

        raise HTTPException(
            status_code=400,

            detail=(
                "date_from cannot be "
                "later than date_to."
            ),
        )


# =========================================================
# ERROR HANDLER
# =========================================================

def _handle_error(result):

    if (
        isinstance(result, dict)
        and result.get("error")
        == "company_not_found"
    ):

        raise HTTPException(
            status_code=404,
            detail="Company not found.",
        )

    return result


# =========================================================
# QUERY PARAMETER HELPER
# =========================================================

def _query_params(
    date_from: Optional[date],

    date_to: Optional[date],

    transaction_type: str,

    payment_method: str,

    search: Optional[str],
):

    return {
        "date_from": date_from,

        "date_to": date_to,

        "transaction_type":
            transaction_type,

        "payment_method":
            payment_method,

        "search": search,
    }


# =========================================================
# GET LEDGER REPORT
# =========================================================

@router.get(
    "/ledger",
    response_model=LedgerReportResponse,
)
def ledger_report(

    date_from: Optional[date] = Query(
        None,
    ),

    date_to: Optional[date] = Query(
        None,
    ),

    transaction_type: str = Query(
        "all",

        pattern=(
            "^(all|"
            "chit_collection|"
            "loan_taken|"
            "loan_given|"
            "loan_repayment_received|"
            "loan_repayment_paid|"
            "expense|"
            "chit_payout)$"
        ),
    ),

    payment_method: str = Query(
        "all",

        pattern=(
            "^(all|"
            "cash|"
            "upi|"
            "bank_transfer|"
            "cheque|"
            "other)$"
        ),
    ),

    search: Optional[str] = Query(
        None,
    ),

    page: int = Query(
        1,
        ge=1,
    ),

    page_size: int = Query(
        10,
        ge=1,
        le=100,
    ),

    db: Session = Depends(
        get_db,
    ),

    current_user: User = Depends(
        get_current_user,
    ),
):

    # =====================================================
    # COMPANY ACCESS
    # =====================================================

    if not current_user.company_id:

        raise HTTPException(
            status_code=403,

            detail=(
                "User is not associated "
                "with a company."
            ),
        )

    # =====================================================
    # DATE VALIDATION
    # =====================================================

    _validate_dates(
        date_from,
        date_to,
    )

    # =====================================================
    # GET REPORT
    # =====================================================

    result = get_ledger_report(

        db=db,

        company_id=(
            current_user.company_id
        ),

        page=page,

        page_size=page_size,

        **_query_params(
            date_from,

            date_to,

            transaction_type,

            payment_method,

            search,
        ),
    )

    return _handle_error(
        result
    )


# =========================================================
# EXCEL EXPORT
# =========================================================

def _export_excel(data):

    workbook = Workbook()

    worksheet = workbook.active

    worksheet.title = "Ledger Report"

    # =====================================================
    # COMPANY HEADER
    # =====================================================

    worksheet["A1"] = (
        data["company_name"]
    )

    worksheet["A1"].font = Font(
        size=16,
        bold=True,
    )

    worksheet["A2"] = (
        f"Company ID: "
        f"{data['company_id']}"
    )

    worksheet["A3"] = (
        "LEDGER REPORT"
    )

    worksheet["A3"].font = Font(
        size=13,
        bold=True,
    )

    # =====================================================
    # SUMMARY
    # =====================================================

    summary = data["summary"]

    summary_rows = [

        (
            "Total Transactions",
            summary[
                "total_transactions"
            ],
        ),

        (
            "Money In",
            summary[
                "total_money_in"
            ],
        ),

        (
            "Money Out",
            summary[
                "total_money_out"
            ],
        ),

        (
            "Net Movement",
            summary[
                "net_movement"
            ],
        ),

        (
            "Cash In",
            summary[
                "cash_in"
            ],
        ),

        (
            "Cash Out",
            summary[
                "cash_out"
            ],
        ),

        (
            "Cash Net",
            summary[
                "cash_net"
            ],
        ),

        (
            "Bank / UPI In",
            summary[
                "bank_in"
            ],
        ),

        (
            "Bank / UPI Out",
            summary[
                "bank_out"
            ],
        ),

        (
            "Bank / UPI Net",
            summary[
                "bank_net"
            ],
        ),
    ]

    row = 5

    for (
        label,
        value,
    ) in summary_rows:

        worksheet.cell(
            row=row,
            column=1,
            value=label,
        )

        worksheet.cell(
            row=row,
            column=2,
            value=value,
        )

        row += 1

    # =====================================================
    # TABLE
    # =====================================================

    row += 1

    headers = [

        "Date",

        "Type",

        "Description",

        "Person",

        "Chit",

        "Payment Method",

        "Flow",

        "Amount",

        "Cash",

        "Bank / UPI",

        "Cash Balance",

        "Bank / UPI Balance",

        "Total Balance",

        "Reference",
    ]

    for column, header in enumerate(
        headers,
        1,
    ):

        cell = worksheet.cell(
            row=row,

            column=column,

            value=header,
        )

        cell.font = Font(
            bold=True,
        )

        cell.alignment = Alignment(
            horizontal="center",
        )

    # =====================================================
    # DATA
    # =====================================================

    for item in data["items"]:

        row += 1

        values = [

            item["date"],

            item[
                "transaction_type"
            ],

            item[
                "description"
            ],

            item.get(
                "person_name"
            ),

            item.get(
                "chit_name"
            ),

            item[
                "payment_method"
            ],

            item[
                "money_flow"
            ],

            item[
                "amount"
            ],

            item[
                "cash_amount"
            ],

            item[
                "bank_amount"
            ],

            item[
                "cash_balance"
            ],

            item[
                "bank_balance"
            ],

            item[
                "total_balance"
            ],

            item.get(
                "reference"
            ),
        ]

        for column, value in enumerate(
            values,
            1,
        ):

            worksheet.cell(
                row=row,

                column=column,

                value=value,
            )

    # =====================================================
    # COLUMN WIDTHS
    # =====================================================

    widths = {

        "A": 14,

        "B": 28,

        "C": 35,

        "D": 22,

        "E": 22,

        "F": 18,

        "G": 12,

        "H": 16,

        "I": 16,

        "J": 18,

        "K": 18,

        "L": 20,

        "M": 18,

        "N": 24,
    }

    for column, width in widths.items():

        worksheet.column_dimensions[
            column
        ].width = width

    # =====================================================
    # OUTPUT
    # =====================================================

    output = BytesIO()

    workbook.save(
        output
    )

    output.seek(0)

    return output


@router.get(
    "/ledger/export/excel"
)
def export_ledger_excel(

    date_from: Optional[date] = Query(
        None,
    ),

    date_to: Optional[date] = Query(
        None,
    ),

    transaction_type: str = Query(
        "all",

        pattern=(
            "^(all|"
            "chit_collection|"
            "loan_taken|"
            "loan_given|"
            "loan_repayment_received|"
            "loan_repayment_paid|"
            "expense|"
            "chit_payout)$"
        ),
    ),

    payment_method: str = Query(
        "all",

        pattern=(
            "^(all|"
            "cash|"
            "upi|"
            "bank_transfer|"
            "cheque|"
            "other)$"
        ),
    ),

    search: Optional[str] = Query(
        None,
    ),

    db: Session = Depends(
        get_db,
    ),

    current_user: User = Depends(
        get_current_user,
    ),
):

    if not current_user.company_id:

        raise HTTPException(
            status_code=403,

            detail=(
                "User is not associated "
                "with a company."
            ),
        )

    _validate_dates(
        date_from,
        date_to,
    )

    data = (
        get_ledger_report_export_data(

            db=db,

            company_id=(
                current_user.company_id
            ),

            **_query_params(
                date_from,

                date_to,

                transaction_type,

                payment_method,

                search,
            ),
        )
    )

    data = _handle_error(
        data
    )

    output = _export_excel(
        data
    )

    return StreamingResponse(

        output,

        media_type=(
            "application/"
            "vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),

        headers={
            "Content-Disposition":
                'attachment; '
                'filename="ledger-report.xlsx"'
        },
    )


# =========================================================
# PDF EXPORT
# =========================================================

@router.get(
    "/ledger/export/pdf"
)
def export_ledger_pdf(

    date_from: Optional[date] = Query(
        None,
    ),

    date_to: Optional[date] = Query(
        None,
    ),

    transaction_type: str = Query(
        "all",

        pattern=(
            "^(all|"
            "chit_collection|"
            "loan_taken|"
            "loan_given|"
            "loan_repayment_received|"
            "loan_repayment_paid|"
            "expense|"
            "chit_payout)$"
        ),
    ),

    payment_method: str = Query(
        "all",

        pattern=(
            "^(all|"
            "cash|"
            "upi|"
            "bank_transfer|"
            "cheque|"
            "other)$"
        ),
    ),

    search: Optional[str] = Query(
        None,
    ),

    db: Session = Depends(
        get_db,
    ),

    current_user: User = Depends(
        get_current_user,
    ),
):

    if not current_user.company_id:

        raise HTTPException(
            status_code=403,

            detail=(
                "User is not associated "
                "with a company."
            ),
        )

    _validate_dates(
        date_from,
        date_to,
    )

    data = (
        get_ledger_report_export_data(

            db=db,

            company_id=(
                current_user.company_id
            ),

            **_query_params(
                date_from,

                date_to,

                transaction_type,

                payment_method,

                search,
            ),
        )
    )

    data = _handle_error(
        data
    )

    # =====================================================
    # PDF DOCUMENT
    # =====================================================

    output = BytesIO()

    document = SimpleDocTemplate(

        output,

        pagesize=landscape(A4),

        rightMargin=24,

        leftMargin=24,

        topMargin=24,

        bottomMargin=24,
    )

    styles = (
        getSampleStyleSheet()
    )

    story = [

        Paragraph(
            f"<b>"
            f"{data['company_name']}"
            f"</b>",

            styles["Title"],
        ),

        Paragraph(
            f"Company ID: "
            f"{data['company_id']}",

            styles["Normal"],
        ),

        Spacer(
            1,
            8,
        ),

        Paragraph(
            "<b>LEDGER REPORT</b>",

            styles["Heading2"],
        ),

        Spacer(
            1,
            8,
        ),
    ]

    # =====================================================
    # SUMMARY
    # =====================================================

    summary = data["summary"]

    story.append(

        Paragraph(

            "Money In: ₹%s "
            "&nbsp;&nbsp; "
            "Money Out: ₹%s "
            "&nbsp;&nbsp; "
            "Net Movement: ₹%s "
            "&nbsp;&nbsp; "
            "Cash Net: ₹%s "
            "&nbsp;&nbsp; "
            "Bank / UPI Net: ₹%s"

            % (

                summary[
                    "total_money_in"
                ],

                summary[
                    "total_money_out"
                ],

                summary[
                    "net_movement"
                ],

                summary[
                    "cash_net"
                ],

                summary[
                    "bank_net"
                ],
            ),

            styles["Normal"],
        )
    )

    story.append(
        Spacer(
            1,
            10,
        )
    )

    # =====================================================
    # TABLE
    # =====================================================

    table_data = [[

        "Date",

        "Type",

        "Description",

        "Person",

        "Chit",

        "Method",

        "Flow",

        "Amount",

        "Cash",

        "Bank/UPI",

        "Cash Bal.",

        "Bank/UPI Bal.",

        "Total Bal.",
    ]]

    for item in data["items"]:

        table_data.append([

            str(
                item["date"]
            ),

            item[
                "transaction_type"
            ],

            item[
                "description"
            ],

            item.get(
                "person_name"
            ) or "",

            item.get(
                "chit_name"
            ) or "",

            item[
                "payment_method"
            ],

            item[
                "money_flow"
            ],

            (
                f"₹"
                f"{item['amount']:,}"
            ),

            (
                f"₹"
                f"{item['cash_amount']:,}"
            ),

            (
                f"₹"
                f"{item['bank_amount']:,}"
            ),

            (
                f"₹"
                f"{item['cash_balance']:,}"
            ),

            (
                f"₹"
                f"{item['bank_balance']:,}"
            ),

            (
                f"₹"
                f"{item['total_balance']:,}"
            ),
        ])

    table = Table(
        table_data,
        repeatRows=1,
    )

    table.setStyle(

        TableStyle([

            (
                "BACKGROUND",
                (0, 0),
                (-1, 0),
                colors.lightgrey,
            ),

            (
                "FONTNAME",
                (0, 0),
                (-1, 0),
                "Helvetica-Bold",
            ),

            (
                "GRID",
                (0, 0),
                (-1, -1),
                0.35,
                colors.grey,
            ),

            (
                "FONTSIZE",
                (0, 0),
                (-1, -1),
                7,
            ),

            (
                "VALIGN",
                (0, 0),
                (-1, -1),
                "MIDDLE",
            ),
        ])
    )

    story.append(
        table
    )

    # =====================================================
    # BUILD PDF
    # =====================================================

    document.build(
        story
    )

    output.seek(0)

    return StreamingResponse(

        output,

        media_type="application/pdf",

        headers={
            "Content-Disposition":
                'attachment; '
                'filename="ledger-report.pdf"'
        },
    )