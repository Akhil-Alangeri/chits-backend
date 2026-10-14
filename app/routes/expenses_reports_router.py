from datetime import date
from io import BytesIO
from typing import Optional

from fastapi import (
    APIRouter,
    Depends,
    HTTPException,
    Query,
    status,
)

from fastapi.responses import StreamingResponse

from openpyxl import Workbook
from openpyxl.styles import (
    Font,
    Alignment,
)

from reportlab.lib import colors
from reportlab.lib.pagesizes import (
    A4,
    landscape,
)

from reportlab.lib.styles import (
    getSampleStyleSheet,
)

from reportlab.platypus import (
    SimpleDocTemplate,
    Table,
    TableStyle,
    Paragraph,
    Spacer,
)

from sqlalchemy.orm import Session

from app.core.database import get_db
from app.core.dependencies import get_current_user

from app.models.admin.user import User

from app.schemas.expenses_reports import (
    ExpenseReportResponse,
)

from app.services.expenses_reports_service import (
    get_expense_report,
    get_expense_report_export_data,
)


from app.services.entitlement_service import require_company_feature

router = APIRouter(
    dependencies=[Depends(require_company_feature("expenses"))],
    prefix="/reports",
    tags=["Reports"],
)


# =========================================================
# HELPERS
# =========================================================

def _validate_date_range(
    date_from: Optional[date],
    date_to: Optional[date],
):
    if (
        date_from
        and date_to
        and date_from > date_to
    ):
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=(
                "date_from cannot be later "
                "than date_to."
            ),
        )


def _handle_company_error(
    exc: ValueError,
):
    if str(exc) == "Company not found.":

        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Company not found.",
        )

    raise exc


# =========================================================
# EXPENSE REPORT
# =========================================================

@router.get(
    "/expenses",
    response_model=ExpenseReportResponse,
)
def expense_report(
    date_from: Optional[date] = Query(
        None
    ),

    date_to: Optional[date] = Query(
        None
    ),

    category: Optional[str] = Query(
        None
    ),

    payment_method: Optional[str] = Query(
        None
    ),

    search: Optional[str] = Query(
        None
    ),

    page: int = Query(
        1,
        ge=1,
    ),

    page_size: int = Query(
        10,
        ge=1,
        le=100,
    ),

    db: Session = Depends(
        get_db
    ),

    current_user: User = Depends(
        get_current_user
    ),
):

    # =====================================================
    # COMPANY
    # =====================================================

    if not current_user.company_id:

        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail=(
                "User is not associated "
                "with a company."
            ),
        )

    # =====================================================
    # DATE VALIDATION
    # =====================================================

    _validate_date_range(
        date_from,
        date_to,
    )

    # =====================================================
    # REPORT
    # =====================================================

    try:

        return get_expense_report(
            db=db,

            company_id=(
                current_user.company_id
            ),

            date_from=date_from,

            date_to=date_to,

            category=category,

            payment_method=(
                payment_method
            ),

            search=search,

            page=page,

            page_size=page_size,
        )

    except ValueError as exc:

        _handle_company_error(
            exc
        )


# =========================================================
# EXCEL EXPORT
# =========================================================

@router.get(
    "/expenses/export/excel"
)
def export_expense_report_excel(
    date_from: Optional[date] = Query(
        None
    ),

    date_to: Optional[date] = Query(
        None
    ),

    category: Optional[str] = Query(
        None
    ),

    payment_method: Optional[str] = Query(
        None
    ),

    search: Optional[str] = Query(
        None
    ),

    db: Session = Depends(
        get_db
    ),

    current_user: User = Depends(
        get_current_user
    ),
):

    if not current_user.company_id:

        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail=(
                "User is not associated "
                "with a company."
            ),
        )

    _validate_date_range(
        date_from,
        date_to,
    )

    try:

        report = (
            get_expense_report_export_data(
                db=db,

                company_id=(
                    current_user.company_id
                ),

                date_from=date_from,

                date_to=date_to,

                category=category,

                payment_method=(
                    payment_method
                ),

                search=search,
            )
        )

    except ValueError as exc:

        _handle_company_error(
            exc
        )

    # =====================================================
    # WORKBOOK
    # =====================================================

    workbook = Workbook()

    worksheet = workbook.active

    worksheet.title = (
        "Expense Report"
    )

    # =====================================================
    # COMPANY HEADER
    # =====================================================

    worksheet["A1"] = (
        report["company_name"]
    )

    worksheet["A1"].font = Font(
        bold=True,
        size=16,
    )

    worksheet["A2"] = (
        f"Company ID: "
        f"{report['company_id']}"
    )

    worksheet["A4"] = (
        "EXPENSE REPORT"
    )

    worksheet["A4"].font = Font(
        bold=True,
        size=14,
    )

    # =====================================================
    # FILTER RANGE
    # =====================================================

    date_range_text = "All Dates"

    if (
        report["date_from"]
        and report["date_to"]
    ):

        date_range_text = (
            f"{report['date_from']} "
            f"to "
            f"{report['date_to']}"
        )

    elif report["date_from"]:

        date_range_text = (
            f"From "
            f"{report['date_from']}"
        )

    elif report["date_to"]:

        date_range_text = (
            f"Up to "
            f"{report['date_to']}"
        )

    worksheet["A5"] = (
        f"Date Range: "
        f"{date_range_text}"
    )

    # =====================================================
    # SUMMARY
    # =====================================================

    summary = report[
        "summary"
    ]

    worksheet["A7"] = (
        "Total Expenses"
    )

    worksheet["B7"] = (
        summary[
            "total_expenses"
        ]
    )

    worksheet["A8"] = (
        "Total Records"
    )

    worksheet["B8"] = (
        summary[
            "total_records"
        ]
    )

    worksheet["A9"] = (
        "Cash Expenses"
    )

    worksheet["B9"] = (
        summary[
            "cash_expenses"
        ]
    )

    worksheet["A10"] = (
        "UPI Expenses"
    )

    worksheet["B10"] = (
        summary[
            "upi_expenses"
        ]
    )

    worksheet["A11"] = (
        "Other Expenses"
    )

    worksheet["B11"] = (
        summary[
            "other_expenses"
        ]
    )

    worksheet["A12"] = (
        "Average Expense"
    )

    worksheet["B12"] = (
        summary[
            "average_expense"
        ]
    )

    worksheet["A13"] = (
        "Highest Expense"
    )

    worksheet["B13"] = (
        summary[
            "highest_expense"
        ]
    )

    worksheet["A14"] = (
        "Categories"
    )

    worksheet["B14"] = (
        summary[
            "unique_categories"
        ]
    )

    # =====================================================
    # CATEGORY SUMMARY
    # =====================================================

    category_start_row = 7

    worksheet["D7"] = (
        "CATEGORY SUMMARY"
    )

    worksheet["D7"].font = Font(
        bold=True
    )

    worksheet["D8"] = "Category"
    worksheet["E8"] = "Amount"
    worksheet["F8"] = "Records"

    for cell in (
        worksheet["D8"],
        worksheet["E8"],
        worksheet["F8"],
    ):

        cell.font = Font(
            bold=True
        )

    category_row = 9

    for item in report[
        "categories"
    ]:

        worksheet.cell(
            row=category_row,
            column=4,
            value=item[
                "category"
            ],
        )

        worksheet.cell(
            row=category_row,
            column=5,
            value=item[
                "amount"
            ],
        )

        worksheet.cell(
            row=category_row,
            column=6,
            value=item[
                "record_count"
            ],
        )

        category_row += 1

    # =====================================================
    # MONTHLY SUMMARY
    # =====================================================

    monthly_start_row = max(
        category_row + 2,
        18,
    )

    worksheet.cell(
        row=monthly_start_row,
        column=1,
        value="MONTHLY SUMMARY",
    )

    worksheet.cell(
        row=monthly_start_row,
        column=1,
    ).font = Font(
        bold=True
    )

    monthly_header_row = (
        monthly_start_row + 1
    )

    monthly_headers = [
        "Year",
        "Month",
        "Amount",
        "Records",
    ]

    for index, header in enumerate(
        monthly_headers,
        start=1,
    ):

        cell = worksheet.cell(
            row=monthly_header_row,
            column=index,
        )

        cell.value = header
        cell.font = Font(
            bold=True
        )

    monthly_row = (
        monthly_header_row + 1
    )

    for item in report[
        "monthly_summary"
    ]:

        values = [
            item["year"],
            item["month_name"],
            item["amount"],
            item["record_count"],
        ]

        for index, value in enumerate(
            values,
            start=1,
        ):

            worksheet.cell(
                row=monthly_row,
                column=index,
                value=value,
            )

        monthly_row += 1

    # =====================================================
    # DETAIL TABLE
    # =====================================================

    detail_start_row = (
        monthly_row + 2
    )

    worksheet.cell(
        row=detail_start_row,
        column=1,
        value="EXPENSE DETAILS",
    )

    worksheet.cell(
        row=detail_start_row,
        column=1,
    ).font = Font(
        bold=True
    )

    header_row = (
        detail_start_row + 1
    )

    headers = [
        "ID",
        "Date",
        "Category",
        "Amount",
        "Payment Method",
        "Description",
        "Reference",
        "Notes",
        "Created By",
        "Created At",
        "Updated At",
    ]

    for index, header in enumerate(
        headers,
        start=1,
    ):

        cell = worksheet.cell(
            row=header_row,
            column=index,
        )

        cell.value = header

        cell.font = Font(
            bold=True
        )

        cell.alignment = Alignment(
            horizontal="center"
        )

    # =====================================================
    # DETAIL ROWS
    # =====================================================

    row_index = (
        header_row + 1
    )

    for item in report[
        "items"
    ]:

        values = [
            item["id"],

            (
                item["expense_date"]
                .isoformat()
                if item[
                    "expense_date"
                ]
                else ""
            ),

            item["category"],

            item["amount"],

            item[
                "payment_method"
            ],

            item[
                "description"
            ]
            or "",

            item[
                "reference"
            ]
            or "",

            item[
                "notes"
            ]
            or "",

            item[
                "created_by"
            ],

            item[
                "created_at"
            ]
            or "",

            item[
                "updated_at"
            ]
            or "",
        ]

        for index, value in enumerate(
            values,
            start=1,
        ):

            worksheet.cell(
                row=row_index,
                column=index,
                value=value,
            )

        row_index += 1

    # =====================================================
    # COLUMN WIDTHS
    # =====================================================

    widths = {
        "A": 10,
        "B": 14,
        "C": 22,
        "D": 16,
        "E": 18,
        "F": 35,
        "G": 22,
        "H": 35,
        "I": 14,
        "J": 22,
        "K": 22,
    }

    for column, width in (
        widths.items()
    ):

        worksheet.column_dimensions[
            column
        ].width = width

    # =====================================================
    # FREEZE
    # =====================================================

    worksheet.freeze_panes = (
        f"A{header_row + 1}"
    )

    # =====================================================
    # OUTPUT
    # =====================================================

    output = BytesIO()

    workbook.save(
        output
    )

    output.seek(0)

    return StreamingResponse(
        output,
        media_type=(
            "application/vnd.openxmlformats-officedocument"
            ".spreadsheetml.sheet"
        ),
        headers={
            "Content-Disposition": (
                'attachment; '
                'filename="expense-report.xlsx"'
            )
        },
    )


# =========================================================
# PDF EXPORT
# =========================================================

@router.get(
    "/expenses/export/pdf"
)
def export_expense_report_pdf(
    date_from: Optional[date] = Query(
        None
    ),

    date_to: Optional[date] = Query(
        None
    ),

    category: Optional[str] = Query(
        None
    ),

    payment_method: Optional[str] = Query(
        None
    ),

    search: Optional[str] = Query(
        None
    ),

    db: Session = Depends(
        get_db
    ),

    current_user: User = Depends(
        get_current_user
    ),
):

    if not current_user.company_id:

        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail=(
                "User is not associated "
                "with a company."
            ),
        )

    _validate_date_range(
        date_from,
        date_to,
    )

    try:

        report = (
            get_expense_report_export_data(
                db=db,

                company_id=(
                    current_user.company_id
                ),

                date_from=date_from,

                date_to=date_to,

                category=category,

                payment_method=(
                    payment_method
                ),

                search=search,
            )
        )

    except ValueError as exc:

        _handle_company_error(
            exc
        )

    # =====================================================
    # PDF
    # =====================================================

    output = BytesIO()

    document = SimpleDocTemplate(
        output,

        pagesize=landscape(A4),

        rightMargin=20,
        leftMargin=20,

        topMargin=20,
        bottomMargin=20,
    )

    styles = (
        getSampleStyleSheet()
    )

    story = []

    # =====================================================
    # COMPANY HEADER
    # =====================================================

    story.append(
        Paragraph(
            report[
                "company_name"
            ],
            styles["Title"],
        )
    )

    story.append(
        Paragraph(
            (
                f"Company ID: "
                f"{report['company_id']}"
            ),
            styles["Normal"],
        )
    )

    story.append(
        Spacer(1, 6)
    )

    story.append(
        Paragraph(
            "EXPENSE REPORT",
            styles["Heading2"],
        )
    )

    # =====================================================
    # DATE RANGE
    # =====================================================

    if (
        report["date_from"]
        and report["date_to"]
    ):

        date_range_text = (
            f"{report['date_from']} "
            f"to "
            f"{report['date_to']}"
        )

    elif report["date_from"]:

        date_range_text = (
            f"From "
            f"{report['date_from']}"
        )

    elif report["date_to"]:

        date_range_text = (
            f"Up to "
            f"{report['date_to']}"
        )

    else:

        date_range_text = (
            "All Dates"
        )

    story.append(
        Paragraph(
            (
                f"Date Range: "
                f"{date_range_text}"
            ),
            styles["Normal"],
        )
    )

    story.append(
        Spacer(1, 10)
    )

    # =====================================================
    # SUMMARY
    # =====================================================

    summary = report[
        "summary"
    ]

    summary_data = [
        [
            "Total",
            "Records",
            "Cash",
            "UPI",
            "Other",
            "Average",
            "Highest",
            "Categories",
        ],

        [
            str(
                summary[
                    "total_expenses"
                ]
            ),

            str(
                summary[
                    "total_records"
                ]
            ),

            str(
                summary[
                    "cash_expenses"
                ]
            ),

            str(
                summary[
                    "upi_expenses"
                ]
            ),

            str(
                summary[
                    "other_expenses"
                ]
            ),

            str(
                summary[
                    "average_expense"
                ]
            ),

            str(
                summary[
                    "highest_expense"
                ]
            ),

            str(
                summary[
                    "unique_categories"
                ]
            ),
        ],
    ]

    summary_table = Table(
        summary_data,
        repeatRows=1,
    )

    summary_table.setStyle(
        TableStyle(
            [
                (
                    "BACKGROUND",
                    (0, 0),
                    (-1, 0),
                    colors.lightgrey,
                ),

                (
                    "FONTNAME",
                    (0, 0),
                    (-1, 0),
                    "Helvetica-Bold",
                ),

                (
                    "GRID",
                    (0, 0),
                    (-1, -1),
                    0.5,
                    colors.grey,
                ),

                (
                    "ALIGN",
                    (0, 0),
                    (-1, -1),
                    "CENTER",
                ),

                (
                    "FONTSIZE",
                    (0, 0),
                    (-1, -1),
                    7,
                ),
            ]
        )
    )

    story.append(
        summary_table
    )

    story.append(
        Spacer(1, 12)
    )

    # =====================================================
    # CATEGORY SUMMARY
    # =====================================================

    story.append(
        Paragraph(
            "Category Summary",
            styles["Heading3"],
        )
    )

    category_data = [
        [
            "Category",
            "Amount",
            "Records",
        ]
    ]

    for item in report[
        "categories"
    ]:

        category_data.append(
            [
                str(
                    item[
                        "category"
                    ]
                ),

                str(
                    item[
                        "amount"
                    ]
                ),

                str(
                    item[
                        "record_count"
                    ]
                ),
            ]
        )

    category_table = Table(
        category_data,
        repeatRows=1,
        colWidths=[
            180,
            100,
            80,
        ],
    )

    category_table.setStyle(
        TableStyle(
            [
                (
                    "BACKGROUND",
                    (0, 0),
                    (-1, 0),
                    colors.lightgrey,
                ),

                (
                    "FONTNAME",
                    (0, 0),
                    (-1, 0),
                    "Helvetica-Bold",
                ),

                (
                    "GRID",
                    (0, 0),
                    (-1, -1),
                    0.4,
                    colors.grey,
                ),

                (
                    "ALIGN",
                    (1, 1),
                    (-1, -1),
                    "CENTER",
                ),

                (
                    "FONTSIZE",
                    (0, 0),
                    (-1, -1),
                    7,
                ),
            ]
        )
    )

    story.append(
        category_table
    )

    story.append(
        Spacer(1, 12)
    )

    # =====================================================
    # MONTHLY SUMMARY
    # =====================================================

    story.append(
        Paragraph(
            "Monthly Summary",
            styles["Heading3"],
        )
    )

    monthly_data = [
        [
            "Year",
            "Month",
            "Amount",
            "Records",
        ]
    ]

    for item in report[
        "monthly_summary"
    ]:

        monthly_data.append(
            [
                str(
                    item["year"]
                ),

                str(
                    item["month_name"]
                ),

                str(
                    item["amount"]
                ),

                str(
                    item[
                        "record_count"
                    ]
                ),
            ]
        )

    monthly_table = Table(
        monthly_data,
        repeatRows=1,
        colWidths=[
            60,
            100,
            100,
            80,
        ],
    )

    monthly_table.setStyle(
        TableStyle(
            [
                (
                    "BACKGROUND",
                    (0, 0),
                    (-1, 0),
                    colors.lightgrey,
                ),

                (
                    "FONTNAME",
                    (0, 0),
                    (-1, 0),
                    "Helvetica-Bold",
                ),

                (
                    "GRID",
                    (0, 0),
                    (-1, -1),
                    0.4,
                    colors.grey,
                ),

                (
                    "ALIGN",
                    (0, 0),
                    (-1, -1),
                    "CENTER",
                ),

                (
                    "FONTSIZE",
                    (0, 0),
                    (-1, -1),
                    7,
                ),
            ]
        )
    )

    story.append(
        monthly_table
    )

    story.append(
        Spacer(1, 14)
    )

    # =====================================================
    # DETAILS
    # =====================================================

    story.append(
        Paragraph(
            "Expense Details",
            styles["Heading3"],
        )
    )

    normal_style = styles[
        "Normal"
    ]

    detail_data = [
        [
            "ID",
            "Date",
            "Category",
            "Amount",
            "Method",
            "Description",
            "Reference",
        ]
    ]

    for item in report[
        "items"
    ]:

        detail_data.append(
            [
                str(
                    item["id"]
                ),

                (
                    item[
                        "expense_date"
                    ].isoformat()
                    if item[
                        "expense_date"
                    ]
                    else ""
                ),

                Paragraph(
                    str(
                        item[
                            "category"
                        ]
                    ),
                    normal_style,
                ),

                str(
                    item["amount"]
                ),

                str(
                    item[
                        "payment_method"
                    ]
                ),

                Paragraph(
                    str(
                        item[
                            "description"
                        ]
                        or ""
                    ),
                    normal_style,
                ),

                Paragraph(
                    str(
                        item[
                            "reference"
                        ]
                        or ""
                    ),
                    normal_style,
                ),
            ]
        )

    detail_table = Table(
        detail_data,
        repeatRows=1,
        colWidths=[
            35,
            65,
            100,
            70,
            65,
            180,
            100,
        ],
    )

    detail_table.setStyle(
        TableStyle(
            [
                (
                    "BACKGROUND",
                    (0, 0),
                    (-1, 0),
                    colors.lightgrey,
                ),

                (
                    "FONTNAME",
                    (0, 0),
                    (-1, 0),
                    "Helvetica-Bold",
                ),

                (
                    "GRID",
                    (0, 0),
                    (-1, -1),
                    0.35,
                    colors.grey,
                ),

                (
                    "VALIGN",
                    (0, 0),
                    (-1, -1),
                    "MIDDLE",
                ),

                (
                    "ALIGN",
                    (0, 0),
                    (-1, -1),
                    "CENTER",
                ),

                (
                    "ALIGN",
                    (2, 1),
                    (2, -1),
                    "LEFT",
                ),

                (
                    "ALIGN",
                    (5, 1),
                    (6, -1),
                    "LEFT",
                ),

                (
                    "FONTSIZE",
                    (0, 0),
                    (-1, -1),
                    6.5,
                ),
            ]
        )
    )

    story.append(
        detail_table
    )

    # =====================================================
    # BUILD
    # =====================================================

    document.build(
        story
    )

    output.seek(0)

    return StreamingResponse(
        output,
        media_type="application/pdf",
        headers={
            "Content-Disposition": (
                'attachment; '
                'filename="expense-report.pdf"'
            )
        },
    )